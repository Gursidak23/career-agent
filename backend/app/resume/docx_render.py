"""Render ResumeData to an ATS-clean .docx (single column, standard headings).

Many ATS parse .docx more reliably than PDF, and this path needs no external
toolchain (unlike Tectonic for PDF). It mirrors the PDF layout: centered name and
headline, sections in the original resume's order, grouped skills and links.
"""

from __future__ import annotations

from pathlib import Path

from ..schemas import ResumeData
from .latex.render import contact_links


def _contact_lines(resume: ResumeData) -> tuple[str, str]:
    c = resume.contact
    first = "  •  ".join(b for b in (c.location, c.phone, c.email) if b)
    second = "  •  ".join(text for _, text in contact_links(resume))
    return first, second


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    """python-docx has no hyperlink API; build the run XML directly."""
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "14487A")
    props.append(color)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def render_docx(resume: ResumeData, out_dir: Path, stem: str) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Inches, Pt

    out_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.5)
        section.top_margin = section.bottom_margin = Inches(0.4)
        if resume.layout.page_size == "letter":
            section.page_width, section.page_height = Inches(8.5), Inches(11)
    right_tab = doc.sections[0].page_width - Inches(1.0)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(0)

    def para(text: str = "", *, bold: bool = False, italic: bool = False):
        p = doc.add_paragraph()
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
        return p

    def two_sided(left: str, right: str, *, bold: bool = False, italic: bool = False) -> None:
        p = para(left, bold=bold, italic=italic)
        if right:
            p.paragraph_format.tab_stops.add_tab_stop(right_tab, WD_TAB_ALIGNMENT.RIGHT)
            r = p.add_run(f"\t{right}")
            r.italic = italic

    def heading(text: str) -> None:
        p = para(text.upper(), bold=True)
        p.paragraph_format.space_before = Pt(6)
        p.runs[0].font.size = Pt(11)

    def bullet(text: str) -> None:
        doc.add_paragraph(text, style="List Bullet")

    def labelled(text: str) -> None:
        label, sep, rest = text.partition(":")
        if sep and 0 < len(label) <= 60 and rest.strip():
            p = para(f"{label}:", bold=True)
            p.add_run(f" {rest.strip()}")
        else:
            para(text)

    # Header
    name = para()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run((resume.contact.name or "").upper())
    run.bold = True
    run.font.size = Pt(20)
    lines = [resume.contact.headline or "", *_contact_lines(resume)]
    for i, line in enumerate(lines):
        if line:
            p = para(line, bold=(i == 0))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    dates = lambda a, b: " - ".join(d for d in (a, b) if d)  # noqa: E731

    for key in resume.ordered_sections():
        if key == "summary":
            heading("Summary")
            para(resume.summary or "")
        elif key == "experience":
            heading("Experience")
            for e in resume.experience:
                two_sided(e.company, e.location or "", bold=True)
                two_sided(e.title, dates(e.start_date, e.end_date), italic=True)
                for b in e.bullets:
                    bullet(b)
        elif key == "projects":
            heading("Projects")
            for proj in resume.projects:
                p = para(proj.name, bold=True)
                if proj.tech:
                    p.add_run(f" | {', '.join(proj.tech)}").italic = True
                if proj.link:
                    p.add_run("  ")
                    _add_hyperlink(p, proj.link, "GitHub" if "github.com" in proj.link else "Link")
                if proj.description:
                    para(proj.description)
                for b in proj.bullets:
                    bullet(b)
        elif key == "skills":
            heading("Technical Skills" if resume.skill_groups else "Skills")
            if resume.skill_groups:
                for g in resume.skill_groups:
                    labelled(f"{g.category}: {', '.join(g.skills)}")
            else:
                para(", ".join(resume.skills))
        elif key == "education":
            heading("Education")
            for ed in resume.education:
                para(ed.degree, bold=True)
                details = [d.strip() for d in (ed.details or "").splitlines() if d.strip()]
                short = [d for d in details if len(d) <= 30]
                para(" | ".join(x for x in [ed.institution, ed.location, dates(ed.start_date, ed.end_date), *short] if x))
                for d in details:
                    if d not in short:
                        labelled(d)
        elif key == "certifications":
            heading("Certifications")
            for c in resume.certifications:
                p = para(" | ".join(x for x in (c.name, c.date) if x))
                if c.url:
                    p.add_run(" | ")
                    _add_hyperlink(p, c.url, "View credential")
        else:
            extra = next(x for x in resume.extra_sections if x.title == key)
            heading(extra.title)
            for item in extra.items:
                labelled(item)

    path = out_dir / f"{stem}.docx"
    doc.save(str(path))
    return path
